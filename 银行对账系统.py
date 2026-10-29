# -*- coding: utf-8 -*-
"""
================================================================
【出纳管理之银行对账系统】 Bank Reconciliation System  （v2 分层版）
================================================================
本项目成果链（对应学习通章节点 3.5.5）：
    ① python 核心代码   -> bank_core.py（对账算法，纯函数，可单测）
    ② AI 协作的 GUI     -> 银行对账系统.py（本文件，Tkinter 桌面界面）
    ③ 交付客户的 exe    -> pyinstaller 打包（见 打包exe.bat / README.md）
    ④ 开源社区          -> git / gitee（见 README.md 上传指南）

v2 相对 v1（单文件版 银行对账系统_单文件版.py）的完善点：
    1. 核心逻辑与界面解耦：对账算法抽到 bank_core.py，18 个单元测试保障；
    2. 对账口径对齐 3.5.5：第一轮【严格匹配】= 结算方式+票号+借方+贷方 四字段全等；
       同时保留两轮【容错匹配】（票号+金额 / 日期+金额），解决实际录入差异，
       界面上用 √（严格）/ ？（容错，建议人工复核）/ 空（未达）三种标记区分；
    3. 期末余额全自动滚动计算（期初 + 逐笔 借-贷），不再需要手工填期末余额；
    4. 新增【保存项目 / 打开项目】：全部数据存为 .json，关掉软件不丢账
       （这也是老师参考版目前缺少、exe 交付客户后必需的能力）；
    5. 新增双击流水行弹窗修改；新增菜单栏；导入导出保留并标注匹配等级。

运行：python3 银行对账系统.py     依赖：openpyxl（pip install openpyxl）
================================================================
"""

import csv
import json
import os
import sys
import datetime
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

import bank_core as bc

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit("缺少 openpyxl，请先执行：pip install openpyxl")


def resource_path(name):
    """定位随程序发布的数据文件：开发时在脚本目录；pyinstaller 单文件模式下在 _MEIPASS 临时目录。"""
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, name)


# ============================================================
# GUI
# ============================================================
class BankReconciliationApp:

    def __init__(self, root):
        self.root = root
        root.title("出纳管理之银行对账系统 v2  （核心算法 bank_core.py ｜ 已通过 18 项单元测试）")
        root.geometry("1380x860")
        root.minsize(1200, 720)

        # ---------- 业务数据 ----------
        self.opening_ent_balance = tk.StringVar(value="602349.00")
        self.opening_bank_balance = tk.StringVar(value="616149.00")
        self.opening = {"ent_open": 602349.0, "bank_open": 616149.0,
                        "ent_unreach": [], "bank_unreach": []}   # 期初未达账明细
        self.records = []          # 全部流水（含期初结转）：FlowRecord
        self.result = None         # bank_core.run_reconciliation 的返回
        self.project_file = None   # 当前项目 .json 路径

        self._build_menu()
        self._build_ui()
        self._refresh_trees()

    # --------------------------------------------------------
    # 菜单栏：文件（新建/打开/保存项目） 帮助
    # --------------------------------------------------------
    def _build_menu(self):
        menubar = tk.Menu(self.root)
        m_file = tk.Menu(menubar, tearoff=0)
        m_file.add_command(label="新建（清空）", command=self.clear_all)
        m_file.add_separator()
        m_file.add_command(label="打开项目 (.json)…", command=self.open_project)
        m_file.add_command(label="保存项目", command=self.save_project)
        m_file.add_command(label="项目另存为…", command=lambda: self.save_project(as_new=True))
        m_file.add_separator()
        m_file.add_command(label="退出", command=self.root.destroy)
        menubar.add_cascade(label="文件", menu=m_file)

        m_help = tk.Menu(menubar, tearoff=0)
        m_help.add_command(label="对账口径说明", command=self.show_rules)
        m_help.add_command(label="关于", command=lambda: messagebox.showinfo(
            "关于", "出纳管理之银行对账系统 v2\n"
                    "核心算法与界面分离架构\n"
                    "单元测试 18 项全部通过"))
        menubar.add_cascade(label="帮助", menu=m_help)
        self.root.config(menu=menubar)

    def show_rules(self):
        messagebox.showinfo("对账口径说明",
                            "第1轮 严格匹配：结算方式+票号+借方金额+贷方金额 四字段全等\n"
                            "        （与学习通 3.5.5 判定口径完全一致），标记 √\n\n"
                            "第2轮 容错匹配：票号+金额+同方向\n"
                            "        （解决两边日期写法不一致导致漏配），标记 ？\n\n"
                            "第3轮 容错匹配：日期+金额+同方向\n"
                            "        （兜底无票号流水，建议人工复核），标记 ？\n\n"
                            "仍未匹配的流水归入未达账四大类，自动生成交叉口径的\n"
                            "《期末银行存款余额调节表》，两端调整后余额相等即对账平衡。")

    # --------------------------------------------------------
    # UI 构建（布局与 v1 一致）
    # --------------------------------------------------------
    def _build_ui(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass

        # ========== 顶部：余额区 ==========
        top = ttk.LabelFrame(
            self.root,
            text="① 余额调节起点（期初余额可从模板 Excel 载入；期末余额对账时自动滚动计算）",
            padding=8)
        top.pack(fill="x", padx=8, pady=(8, 4))

        ttk.Label(top, text="企业日记账期初余额：").grid(row=0, column=0, sticky="e")
        ttk.Entry(top, textvariable=self.opening_ent_balance, width=13).grid(row=0, column=1, padx=2)
        ttk.Label(top, text="元").grid(row=0, column=2, padx=(0, 14))

        ttk.Label(top, text="银行对账单期初余额：").grid(row=0, column=3, sticky="e")
        ttk.Entry(top, textvariable=self.opening_bank_balance, width=13).grid(row=0, column=4, padx=2)
        ttk.Label(top, text="元").grid(row=0, column=5, padx=(0, 14))

        ttk.Button(top, text="从模板 Excel 载入期初",
                   command=self.load_opening_excel).grid(row=0, column=6, padx=4)
        ttk.Button(top, text="载入示例数据",
                   command=self.load_sample_data).grid(row=0, column=7, padx=4)
        ttk.Button(top, text="保存项目", command=self.save_project).grid(row=0, column=8, padx=4)
        ttk.Button(top, text="打开项目", command=self.open_project).grid(row=0, column=9, padx=4)

        # ========== 中部：流水录入 ==========
        mid = ttk.LabelFrame(self.root, text="② 本期流水录入 / 导入（企业日记账 与 银行对账单）", padding=8)
        mid.pack(fill="x", padx=8, pady=4)

        self.src_var = tk.StringVar(value=bc.SRC_ENTERPRISE)
        ttk.Label(mid, text="数据来源：").grid(row=0, column=0, sticky="e")
        ttk.Radiobutton(mid, text=bc.SRC_ENTERPRISE, variable=self.src_var,
                        value=bc.SRC_ENTERPRISE).grid(row=0, column=1, padx=4)
        ttk.Radiobutton(mid, text=bc.SRC_BANK, variable=self.src_var,
                        value=bc.SRC_BANK).grid(row=0, column=2, padx=4)

        ttk.Label(mid, text="日期：").grid(row=0, column=3, sticky="e")
        self.ent_date = ttk.Entry(mid, width=10)
        self.ent_date.insert(0, datetime.date.today().strftime("%Y-%m-%d"))
        self.ent_date.grid(row=0, column=4, padx=2)

        ttk.Label(mid, text="结算方式：").grid(row=0, column=5, sticky="e")
        self.ent_method = ttk.Entry(mid, width=8)
        self.ent_method.grid(row=0, column=6, padx=2)

        ttk.Label(mid, text="票号：").grid(row=0, column=7, sticky="e")
        self.ent_ticket = ttk.Entry(mid, width=10)
        self.ent_ticket.grid(row=0, column=8, padx=2)

        ttk.Label(mid, text="借方金额：").grid(row=1, column=0, sticky="e", pady=(4, 0))
        self.ent_debit = ttk.Entry(mid, width=11)
        self.ent_debit.grid(row=1, column=1, padx=2, pady=(4, 0))

        ttk.Label(mid, text="贷方金额：").grid(row=1, column=2, sticky="e", pady=(4, 0))
        self.ent_credit = ttk.Entry(mid, width=11)
        self.ent_credit.grid(row=1, column=3, padx=2, pady=(4, 0))

        ttk.Label(mid, text="期末余额：(留空自动滚算)").grid(row=1, column=4, sticky="e", pady=(4, 0))
        self.ent_balance = ttk.Entry(mid, width=11)
        self.ent_balance.grid(row=1, column=5, padx=2, pady=(4, 0))

        ttk.Button(mid, text="添加一笔", command=self.add_record).grid(row=1, column=6, padx=6, pady=(4, 0))
        ttk.Button(mid, text="导入 Excel", command=self.import_excel).grid(row=1, column=7, padx=4, pady=(4, 0))
        ttk.Button(mid, text="导入 CSV/TXT", command=self.import_csv).grid(row=1, column=8, padx=4, pady=(4, 0))
        ttk.Button(mid, text="删除选中行", command=self.delete_selected).grid(row=1, column=9, padx=4, pady=(4, 0))

        hint = ("说明：借方、贷方只填其中一侧。双击流水行可修改。"
                "对账标记图例：√=严格匹配（四字段全等，3.5.5 口径）　"
                "？=容错匹配（票号/日期+金额，建议人工复核）　空=未达账")
        ttk.Label(mid, text=hint, foreground="#666").grid(
            row=2, column=0, columnspan=10, sticky="w", pady=(6, 0))

        # ========== 主体 PanedWindow：上=流水；下左=未达账 下右=余额调节表 ==========
        body = ttk.PanedWindow(self.root, orient=tk.VERTICAL)
        body.pack(fill="both", expand=True, padx=8, pady=4)

        # --- 上：全部流水 ---
        self.top_body = ttk.LabelFrame(body, text="③ 全部流水清单（0 笔）", padding=4)
        body.add(self.top_body, weight=3)

        cols = ("phase", "source", "date", "settlement_method", "ticket_no",
                "debit", "credit", "balance", "matched")
        self.flow_tree = ttk.Treeview(self.top_body, columns=cols, show="headings", height=10)
        headings = {
            "phase": "期间", "source": "数据来源", "date": "日期",
            "settlement_method": "结算方式", "ticket_no": "票号",
            "debit": "借方金额", "credit": "贷方金额", "balance": "滚动余额",
            "matched": "已达",
        }
        widths = {"phase": 55, "source": 105, "date": 100, "settlement_method": 85,
                  "ticket_no": 100, "debit": 115, "credit": 115, "balance": 120, "matched": 55}
        for c in cols:
            self.flow_tree.heading(c, text=headings[c])
            self.flow_tree.column(c, width=widths[c], anchor="center")
        vsb = ttk.Scrollbar(self.top_body, orient="vertical", command=self.flow_tree.yview)
        self.flow_tree.configure(yscrollcommand=vsb.set)
        self.flow_tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")
        self.flow_tree.bind("<Double-1>", self.edit_selected)
        self.flow_tree.tag_configure("loose", background="#fff8e6")   # 容错匹配行淡黄底
        self.flow_tree.tag_configure("opening", foreground="#7a5c00")

        # --- 下：未达账 + 余额调节表 ---
        bottom = ttk.PanedWindow(body, orient=tk.HORIZONTAL)
        body.add(bottom, weight=2)

        self.left = ttk.LabelFrame(bottom, text="④ 未达账项明细（0 笔）", padding=4)
        bottom.add(self.left, weight=1)
        ucols = ("category", "source", "date", "settlement_method", "ticket_no", "amount", "phase")
        self.unreach_tree = ttk.Treeview(self.left, columns=ucols, show="headings", height=8)
        uhead = {"category": "类别", "source": "来源", "date": "日期", "settlement_method": "结算方式",
                 "ticket_no": "票号", "amount": "金额", "phase": "期间"}
        uwidth = {"category": 155, "source": 85, "date": 95, "settlement_method": 75,
                  "ticket_no": 95, "amount": 105, "phase": 50}
        for c in ucols:
            self.unreach_tree.heading(c, text=uhead[c])
            self.unreach_tree.column(c, width=uwidth[c], anchor="center")
        self.unreach_tree.pack(fill="both", expand=True)

        right = ttk.LabelFrame(bottom, text="⑤ 银行存款余额调节表", padding=6)
        bottom.add(right, weight=1)
        self.statement_text = tk.Text(right, height=12, width=60, font=("Consolas", 11))
        self.statement_text.pack(fill="both", expand=True)
        op = ttk.Frame(right)
        op.pack(fill="x", pady=(6, 0))
        ttk.Button(op, text="开始自动对账", command=self.do_reconcile).pack(side="left", padx=4)
        ttk.Button(op, text="导出 Excel (.xlsx)", command=self.export_excel).pack(side="left", padx=4)
        ttk.Button(op, text="导出 CSV", command=self.export_csv).pack(side="left", padx=4)
        ttk.Button(op, text="清空全部", command=self.clear_all).pack(side="left", padx=4)
        self.check_label = ttk.Label(op, text="校验结果：尚未对账", foreground="#1565c0",
                                     font=("", 11, "bold"))
        self.check_label.pack(side="left", padx=16)

    # --------------------------------------------------------
    # 录入 / 编辑 / 删除
    # --------------------------------------------------------
    def _read_amounts(self, debit_txt, credit_txt):
        debit = bc.validate_amount(debit_txt)
        credit = bc.validate_amount(credit_txt)
        if debit is None or credit is None:
            messagebox.showwarning("金额错误", "金额必须是数字")
            return None
        if debit > 0 and credit > 0:
            messagebox.showwarning("金额冲突", "同一笔流水只能填借方或贷方其中之一")
            return None
        if debit == 0 and credit == 0:
            messagebox.showwarning("金额为空", "请填写借方或贷方金额（必须 > 0）")
            return None
        return debit, credit

    def add_record(self):
        date = bc.validate_date(self.ent_date.get())
        if date is None:
            messagebox.showwarning("日期格式错误", "日期应为 YYYY-MM-DD（如 2025-12-08）")
            return
        amounts = self._read_amounts(self.ent_debit.get(), self.ent_credit.get())
        if amounts is None:
            return
        debit, credit = amounts
        balance = bc.validate_amount(self.ent_balance.get()) or 0.0   # 可留空，滚动时重算
        self.records.append(bc.FlowRecord(
            date=date, settlement_method=self.ent_method.get().strip(),
            ticket_no=self.ent_ticket.get().strip(),
            debit=debit, credit=credit, balance=balance,
            source=self.src_var.get()))
        for w in (self.ent_method, self.ent_ticket, self.ent_debit, self.ent_credit, self.ent_balance):
            w.delete(0, tk.END)
        self.result = None
        self._refresh_trees()

    def edit_selected(self, event=None):
        """双击流水行 → 弹窗修改 / 删除该行。"""
        sel = self.flow_tree.selection()
        if not sel:
            return
        idx = int(sel[0])
        if not (0 <= idx < len(self.records)):
            return
        rec = self.records[idx]
        win = tk.Toplevel(self.root)
        win.title("修改第 {} 行（{}）".format(idx + 1, rec.source))
        win.transient(self.root)
        win.grab_set()
        frm = ttk.Frame(win, padding=12)
        frm.pack()
        ents = {}
        fields = [("日期", rec.date), ("结算方式", rec.settlement_method),
                  ("票号", rec.ticket_no),
                  ("借方金额", rec.debit if rec.debit else ""),
                  ("贷方金额", rec.credit if rec.credit else "")]
        for i, (lab, val) in enumerate(fields):
            ttk.Label(frm, text=lab + "：").grid(row=i, column=0, sticky="e", pady=3)
            e = ttk.Entry(frm, width=16)
            e.insert(0, str(val))
            e.grid(row=i, column=1, pady=3, padx=4)
            ents[lab] = e

        def apply():
            d = bc.validate_date(ents["日期"].get())
            if d is None:
                messagebox.showwarning("日期错误", "日期格式应为 YYYY-MM-DD", parent=win)
                return
            amounts = self._read_amounts(ents["借方金额"].get(), ents["贷方金额"].get())
            if amounts is None:
                return
            rec.date = d
            rec.settlement_method = ents["结算方式"].get().strip()
            rec.ticket_no = ents["票号"].get().strip()
            rec.debit, rec.credit = amounts
            rec.matched, rec.match_key = "", None
            self.result = None
            self._refresh_trees()
            win.destroy()

        def remove():
            del self.records[idx]
            self.result = None
            self._refresh_trees()
            win.destroy()

        bb = ttk.Frame(win, padding=(0, 8))
        bb.pack()
        ttk.Button(bb, text="保存修改", command=apply).pack(side="left", padx=6)
        ttk.Button(bb, text="删除此行", command=remove).pack(side="left", padx=6)
        ttk.Button(bb, text="取消", command=win.destroy).pack(side="left", padx=6)

    def delete_selected(self):
        sel = self.flow_tree.selection()
        if not sel:
            messagebox.showinfo("提示", "请先在流水清单中选中要删除的行")
            return
        idxs = sorted([int(i) for i in sel], reverse=True)
        for i in idxs:
            if 0 <= i < len(self.records):
                del self.records[i]
        self.result = None
        self._refresh_trees()

    def clear_all(self):
        if not messagebox.askyesno("确认清空", "确定要清空全部数据，重新开始吗？"):
            return
        self.records.clear()
        self.opening = {"ent_open": 0.0, "bank_open": 0.0, "ent_unreach": [], "bank_unreach": []}
        self.result = None
        self.project_file = None
        self.opening_ent_balance.set("0.00")
        self.opening_bank_balance.set("0.00")
        self.statement_text.delete("1.0", tk.END)
        self.check_label.config(text="校验结果：尚未对账", foreground="#1565c0")
        self._refresh_trees()

    # --------------------------------------------------------
    # 导入：Excel（模板 Sheet1/2/3）与 CSV
    # --------------------------------------------------------
    def _apply_opening(self, opening):
        """把 bank_core 解析出的期初数据落到界面状态。"""
        self.opening = opening
        if opening.get("ent_open"):
            self.opening_ent_balance.set("{:.2f}".format(opening["ent_open"]))
        if opening.get("bank_open"):
            self.opening_bank_balance.set("{:.2f}".format(opening["bank_open"]))
        self._carry_forward_opening()

    def _carry_forward_opening(self):
        """期初未达账明细 → FlowRecord（is_opening=True）并入主流水表（先清除旧结转，避免重复）。"""
        self.records = [r for r in self.records if not r.is_opening]
        self.records.extend(bc.opening_items_to_records(self.opening))

    def import_csv(self):
        path = filedialog.askopenfilename(
            title="选择 CSV / TXT 流水文件",
            filetypes=[("数据文件", "*.csv *.txt"), ("CSV", "*.csv"), ("TXT", "*.txt")])
        if not path:
            return
        try:
            rows = bc.read_flow_csv(path, self.src_var.get())
        except Exception as e:
            messagebox.showerror("导入失败", str(e))
            return
        if not rows:
            messagebox.showwarning("导入失败", "没有解析到有效流水（需要 日期,结算方式,票号,借方,贷方 列）。")
            return
        self.records.extend(rows)
        self.result = None
        self._refresh_trees()
        messagebox.showinfo("导入完成", "成功导入 {} 笔到【{}】。\n（期末余额由系统自动滚动计算）".format(
            len(rows), self.src_var.get()))

    def import_excel(self):
        """从模板 Excel 导入（Sheet1 期初调节表 / Sheet2 日记账 / Sheet3 对账单）。"""
        path = filedialog.askopenfilename(
            title="选择模板 Excel 文件",
            filetypes=[("Excel 文件", "*.xlsx"), ("所有文件", "*.*")])
        if not path:
            return
        try:
            wb = load_workbook(path, data_only=True)
        except Exception as e:
            messagebox.showerror("打开失败", str(e))
            return

        loaded, flow_added = 0, 0
        for sname, source in (("Sheet2", bc.SRC_ENTERPRISE), ("Sheet3", bc.SRC_BANK)):
            if sname not in wb.sheetnames:
                continue
            rows = list(wb[sname].iter_rows(values_only=True))
            added_rows = bc.parse_flow_rows(rows, source)
            if added_rows:
                self.records.extend(added_rows)
                flow_added += len(added_rows)
                loaded += 1

        # Sheet1：期初余额调节表（含期初未达账明细）
        if "Sheet1" in wb.sheetnames:
            rows1 = list(wb["Sheet1"].iter_rows(values_only=True))
            opening = bc.parse_opening_rows(rows1)
            if opening["ent_open"] or opening["bank_open"] or \
                    opening["ent_unreach"] or opening["bank_unreach"]:
                self._apply_opening(opening)
                loaded += 1

        if loaded == 0:
            messagebox.showwarning(
                "未识别到有效数据",
                "请确认 Excel 是否包含 Sheet1（期初余额调节表）、\n"
                "Sheet2（企业日记账）或 Sheet3（银行对账单）。")
            return
        self.result = None
        self._refresh_trees()
        messagebox.showinfo("导入完成",
                            "已载入模板：本期流水 {} 笔，期初未达账已自动结转。\n点击【开始自动对账】继续。".format(flow_added))

    def load_opening_excel(self):
        """仅载入期初余额调节表（Sheet1 或 [节名] 分节 CSV 皆可）。"""
        path = filedialog.askopenfilename(
            title="选择期初余额调节表（.xlsx 或 .csv）",
            filetypes=[("表格文件", "*.xlsx *.csv"), ("所有文件", "*.*")])
        if not path:
            return
        try:
            if path.lower().endswith(".csv"):
                opening = bc.read_opening_csv(path)
            else:
                wb = load_workbook(path, data_only=True)
                if "Sheet1" not in wb.sheetnames:
                    messagebox.showwarning("未找到 Sheet1", "模板 Excel 中应包含 Sheet1（期初余额调节表）")
                    return
                opening = bc.parse_opening_rows(list(wb["Sheet1"].iter_rows(values_only=True)))
        except Exception as e:
            messagebox.showerror("载入失败", str(e))
            return
        self._apply_opening(opening)
        self.result = None
        self._refresh_trees()
        messagebox.showinfo("期初数据已载入",
                            "企业期初余额：{:,.2f} 元\n银行期初余额：{:,.2f} 元\n"
                            "企业方期初未达账 {} 笔\n银行方期初未达账 {} 笔\n（已自动结转进入流水表）".format(
                                self.opening["ent_open"], self.opening["bank_open"],
                                len(self.opening["ent_unreach"]), len(self.opening["bank_unreach"])))

    # --------------------------------------------------------
    # 核心：自动对账（调用 bank_core）
    # --------------------------------------------------------
    def do_reconcile(self):
        try:
            ent_open = float(self.opening_ent_balance.get() or 0)
            bank_open = float(self.opening_bank_balance.get() or 0)
        except ValueError:
            messagebox.showwarning("余额错误", "期初余额必须是数字")
            return
        self.opening["ent_open"], self.opening["bank_open"] = ent_open, bank_open

        if not any(r.source == bc.SRC_ENTERPRISE for r in self.records) or \
                not any(r.source == bc.SRC_BANK for r in self.records):
            messagebox.showwarning(
                "数据不足",
                "请同时录入【企业日记账】和【银行对账单】两类数据后再对账。\n"
                "可点击【载入示例数据】或【导入 Excel / CSV】快速开始。")
            return

        try:
            self.result = bc.run_reconciliation(self.records, ent_open, bank_open)
        except Exception as e:
            messagebox.showerror("对账失败", str(e))
            return
        self._refresh_trees()
        self._generate_statement()

    def _generate_statement(self):
        st = self.result["statement"]
        self.statement_text.delete("1.0", tk.END)
        self.statement_text.insert("1.0", bc.statement_text(st))
        if st["平衡"]:
            self.check_label.config(text="校验结果：✔ 调节后余额相等，对账平衡", foreground="#2e7d32")
        else:
            self.check_label.config(text="校验结果：✘ 调节后余额不相等，请核查", foreground="#c62828")

    # --------------------------------------------------------
    # 项目持久化：保存 / 打开 .json
    # --------------------------------------------------------
    def save_project(self, as_new=False):
        path = self.project_file
        if as_new or not path:
            path = filedialog.asksaveasfilename(
                title="保存对账项目",
                defaultextension=".json",
                initialfile="银行对账项目_" + datetime.datetime.now().strftime("%Y%m%d_%H%M") + ".json",
                filetypes=[("对账项目", "*.json")])
            if not path:
                return
        try:
            self.opening["ent_open"] = bc.to_float(self.opening_ent_balance.get())
            self.opening["bank_open"] = bc.to_float(self.opening_bank_balance.get())
            data = {
                "app": "出纳管理之银行对账系统", "version": 2,
                "saved_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "opening": self.opening,
                "records": [r.to_dict() for r in self.records],
                "result": ({"stats": self.result["stats"], "statement": self.result["statement"]}
                           if self.result else None),
            }
            with open(path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=1)
        except Exception as e:
            messagebox.showerror("保存失败", str(e))
            return
        self.project_file = path
        messagebox.showinfo("保存成功", "项目已保存到：\n{}\n\n下次打开软件选择【打开项目】即可继续对账。".format(path))

    def open_project(self):
        path = filedialog.askopenfilename(
            title="打开对账项目",
            filetypes=[("对账项目", "*.json"), ("所有文件", "*.*")])
        if not path:
            return
        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
            self.opening = data.get("opening", {"ent_open": 0, "bank_open": 0,
                                                "ent_unreach": [], "bank_unreach": []})
            self.records = [bc.FlowRecord.from_dict(d) for d in data.get("records", [])]
            self.opening_ent_balance.set("{:.2f}".format(self.opening.get("ent_open", 0)))
            self.opening_bank_balance.set("{:.2f}".format(self.opening.get("bank_open", 0)))
        except Exception as e:
            messagebox.showerror("打开失败", "不是有效的对账项目文件：\n{}".format(e))
            return
        self.project_file = path
        self.result = None
        self._refresh_trees()
        messagebox.showinfo("打开成功",
                            "已载入 {} 笔流水。\n数据有变动时请点击【开始自动对账】重新计算。".format(len(self.records)))

    # --------------------------------------------------------
    # 载入示例数据（优先读仓库 CSV，缺失时用内置账套——打包 exe 后依然可用）
    # --------------------------------------------------------
    def load_sample_data(self):
        if not messagebox.askyesno("载入示例数据", "将清空当前数据并载入教材示例账套（表3-14/3-15/3-16），继续？"):
            return
        self.records.clear()
        self.result = None
        self.opening = {"ent_open": 602349.0, "bank_open": 616149.0,
                        "ent_unreach": [], "bank_unreach": []}

        used_file = False
        try:
            opening = bc.read_opening_csv(resource_path("示例数据_期初余额调节表.csv"))
            ent = bc.read_flow_csv(resource_path("示例数据_单位银行存款日记账.csv"), bc.SRC_ENTERPRISE)
            bank = bc.read_flow_csv(resource_path("示例数据_银行对账单.csv"), bc.SRC_BANK)
            if ent and bank:
                self.opening = opening
                self.records = bc.opening_items_to_records(opening) + ent + bank
                used_file = True
        except Exception:
            used_file = False
        if not used_file:                       # exe 单文件模式 / CSV 缺失 → 内置账套
            sample = bc.load_builtin_sample()
            self.opening = {"ent_open": sample["ent_open"], "bank_open": sample["bank_open"],
                            "ent_unreach": [], "bank_unreach": []}
            self.records = list(sample["records"])

        self.opening_ent_balance.set("{:.2f}".format(self.opening["ent_open"]))
        self.opening_bank_balance.set("{:.2f}".format(self.opening["bank_open"]))
        self.do_reconcile()
        st = self.result["stats"]
        messagebox.showinfo(
            "示例数据已载入并完成对账" + ("（内置账套）" if not used_file else ""),
            "企业日记账 {} 笔 + 期初结转，银行对账单 {} 笔\n\n"
            "对账结果：\n"
            "  • 严格匹配 {} 对（√）　容错匹配 {} 对（？）\n"
            "  • 企业方未达 {} 笔　银行方未达 {} 笔\n"
            "  • 两端调整后余额 {:,.2f} 元（平衡）".format(
                st["企业方总笔数"], st["银行方总笔数"], st["严格匹配对数"], st["容错匹配对数"],
                st["企业方未达笔数"], st["银行方未达笔数"],
                self.result["statement"]["企业调整后余额"]))

    # --------------------------------------------------------
    # 表格刷新
    # --------------------------------------------------------
    def _refresh_trees(self):
        # 流水表
        for item in self.flow_tree.get_children():
            self.flow_tree.delete(item)
        for idx, r in enumerate(self.records):
            tags = ()
            if r.matched == bc.MATCH_LOOSE:
                tags = ("loose",)
            elif r.is_opening:
                tags = ("opening",)
            self.flow_tree.insert("", "end", iid=str(idx), values=r.tuple_for_tree(), tags=tags)
        self.top_body.config(text="③ 全部流水清单（含期初结转 + 本期录入，共 {} 笔）".format(len(self.records)))

        # 未达账表
        for item in self.unreach_tree.get_children():
            self.unreach_tree.delete(item)
        n_unreach = 0
        if self.result:
            for u in self.result["unreached"]:
                r = u["record"]
                self.unreach_tree.insert("", "end", values=(
                    u["category"], r.source, r.date, r.settlement_method, r.ticket_no,
                    "{:,.2f}".format(r.amount()), "期初" if r.is_opening else "本期"))
            n_unreach = len(self.result["unreached"])
        self.left.config(text="④ 未达账项明细（共 {} 笔）".format(n_unreach))

    # --------------------------------------------------------
    # 导出 Excel（6 个 Sheet，含匹配等级）
    # --------------------------------------------------------
    def export_excel(self):
        if not self.result:
            messagebox.showinfo("提示", "请先点击【开始自动对账】生成结果后再导出。")
            return
        path = filedialog.asksaveasfilename(
            title="保存为 Excel", defaultextension=".xlsx",
            initialfile="银行对账结果_" + datetime.datetime.now().strftime("%Y%m%d_%H%M%S"),
            filetypes=[("Excel 文件", "*.xlsx")])
        if not path:
            return
        try:
            wb = Workbook()
            ws1 = wb.active
            ws1.title = "期初余额调节表"
            self._write_reconciliation_sheet(ws1, mode="opening")
            ws2 = wb.create_sheet("单位银行存款日记账")
            self._write_flow_sheet(ws2, bc.SRC_ENTERPRISE)
            ws3 = wb.create_sheet("银行对账单")
            self._write_flow_sheet(ws3, bc.SRC_BANK)
            ws4 = wb.create_sheet("期末余额调节表")
            self._write_reconciliation_sheet(ws4, mode="ending")
            ws5 = wb.create_sheet("已达账项")
            self._write_reconciled_sheet(ws5)
            ws6 = wb.create_sheet("未达账项明细")
            self._write_unreached_sheet(ws6)
            wb.save(path)
        except Exception as e:
            messagebox.showerror("导出失败", str(e))
            return
        messagebox.showinfo("导出成功", "结果已保存至：\n{}\n\n包含 6 个 Sheet：\n"
                                       "期初余额调节表 / 单位银行存款日记账 / 银行对账单 /\n"
                                       "期末余额调节表 / 已达账项明细 / 未达账项明细".format(path))

    @staticmethod
    def _style_header(ws, row, headers, border, fill):
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=row, column=i, value=h)
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center")
            c.border = border
            c.fill = fill

    def _write_reconciliation_sheet(self, ws, mode="ending"):
        st = self.result["statement"] if self.result else bc.build_statement(
            self.opening["ent_open"], self.opening["bank_open"],
            [], [], [], [])
        title = "期初余额调节表" if mode == "opening" else "期末余额调节表"
        if mode == "opening":
            # 期初口径：直接用期初余额 + 期初未达账
            ent_open, bank_open = self.opening["ent_open"], self.opening["bank_open"]
            b_recv = round(sum(it["debit"] for it in self.opening["bank_unreach"]), 2)
            b_pay = round(sum(it["credit"] for it in self.opening["bank_unreach"]), 2)
            e_recv = round(sum(it["debit"] for it in self.opening["ent_unreach"]), 2)
            e_pay = round(sum(it["credit"] for it in self.opening["ent_unreach"]), 2)
            ent_bal, bank_bal = ent_open, bank_open
        else:
            ent_bal = st["企业期末余额"]
            bank_bal = st["银行期末余额"]
            b_recv, b_pay = st["企业_银行已收企业未收"], st["企业_银行已付企业未付"]
            e_recv, e_pay = st["银行_企业已收银行未收"], st["银行_企业已付银行未付"]
        adj_ent = round(ent_bal + b_recv - b_pay, 2)
        adj_bank = round(bank_bal + e_recv - e_pay, 2)

        ws.merge_cells("A1:E1")
        ws["A1"] = "表3-14  " + title
        ws["A1"].font = Font(name="黑体", size=14, bold=True)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        thin = Side(border_style="thin", color="333333")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        fill = PatternFill("solid", fgColor="E8EEF4")
        self._style_header(ws, 2, ["单位日记账项目", "日记账余额", "银行对账单项目", "对账单余额"], border, fill)

        rows_data = [
            ["调整前余额：", ent_bal, "调整前余额：", bank_bal],
            ["加：银行已收，企业未收", b_recv, "加：企业已收，银行未收", e_recv],
            ["减：银行已付，企业未付", b_pay, "减：企业已付，银行未付", e_pay],
            ["调整后余额：", adj_ent, "调整后余额：", adj_bank],
        ]
        for r, row in enumerate(rows_data, 3):
            for i, val in enumerate(row, 1):
                c = ws.cell(row=r, column=i, value=val)
                c.border = border
                c.alignment = Alignment(horizontal="center" if i in (2, 4) else "left")
                if i in (2, 4) and isinstance(val, (int, float)):
                    c.number_format = "#,##0.00"
                if r == 6:
                    c.font = Font(bold=True)

        # 期末模式：附未达账明细
        if mode == "ending" and self.result:
            headers = ["日期", "结算方式", "票号", "借方金额", "贷方金额"]
            start = 9
            for title_txt, items in (
                    ("期末企业方未达账（企业已记、银行未记）", self.result["ent_only"]),
                    ("期末银行方未达账（银行已记、企业未记）", self.result["bank_only"])):
                start += 1
                ws.cell(row=start, column=1, value="{}：共 {} 笔".format(title_txt, len(items)))\
                    .font = Font(bold=True, color="C00000")
                start += 1
                self._style_header(ws, start, headers, border, fill)
                for k, r0 in enumerate(items, 1):
                    vals = [r0.date, r0.settlement_method, r0.ticket_no,
                            r0.debit or "", r0.credit or ""]
                    for i, v in enumerate(vals, 1):
                        c = ws.cell(row=start + k, column=i, value=v)
                        c.border = border
                        if i in (4, 5) and isinstance(v, (int, float)):
                            c.number_format = "#,##0.00"
                start += len(items)

        for i, w in enumerate([20, 15, 20, 15, 12], 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def _write_flow_sheet(self, ws, source):
        ws.merge_cells("A1:F1")
        ws["A1"] = ("表3-15  单位银行存款日记账" if source == bc.SRC_ENTERPRISE
                    else "表3-16  银行对账单")
        ws["A1"].font = Font(name="黑体", size=13, bold=True)
        ws["A1"].alignment = Alignment(horizontal="center")
        thin = Side(border_style="thin", color="333333")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        fill = PatternFill("solid", fgColor="E8EEF4")
        self._style_header(ws, 2, ["日期", "结算方式", "票号", "借方金额", "贷方金额", "滚动余额", "对账"], border, fill)
        rows = sorted([r for r in self.records if r.source == source and not r.is_opening],
                      key=lambda r: r.date)
        for k, r in enumerate(rows, 3):
            vals = [r.date, r.settlement_method, r.ticket_no,
                    r.debit or None, r.credit or None, r.balance or None, r.mark_label()]
            for i, v in enumerate(vals, 1):
                c = ws.cell(row=k, column=i, value=v)
                c.border = border
                c.alignment = Alignment(horizontal="center" if i != 1 else "left")
                if i in (4, 5, 6) and isinstance(v, (int, float)):
                    c.number_format = "#,##0.00"
        for i, w in enumerate([12, 10, 12, 13, 13, 14, 8], 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def _write_reconciled_sheet(self, ws):
        ws["A1"] = "已达账项（√=严格匹配  ？=容错匹配，建议人工复核）"
        ws["A1"].font = Font(name="黑体", size=13, bold=True)
        thin = Side(border_style="thin", color="333333")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        fill = PatternFill("solid", fgColor="E8EEF4")
        self._style_header(ws, 2, ["企业日期", "企业票号", "企业方向", "金额",
                                   "银行日期", "银行票号", "银行方向", "金额", "匹配等级"], border, fill)
        for k, (er, br, level) in enumerate(self.result["pairs"], 3):
            vals = [er.date, er.ticket_no, er.direction(), er.amount(),
                    br.date, br.ticket_no, br.direction(), br.amount(), level + "匹配"]
            for i, v in enumerate(vals, 1):
                c = ws.cell(row=k, column=i, value=v)
                c.border = border
                if i in (4, 8) and isinstance(v, (int, float)):
                    c.number_format = "#,##0.00"
        for i, w in enumerate([12, 12, 10, 13, 12, 12, 10, 13, 12], 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def _write_unreached_sheet(self, ws):
        ws["A1"] = "未达账项明细（按类别分列，含期初结转）"
        ws["A1"].font = Font(name="黑体", size=13, bold=True)
        thin = Side(border_style="thin", color="333333")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        fill = PatternFill("solid", fgColor="E8EEF4")
        self._style_header(ws, 2, ["类别", "来源", "日期", "结算方式", "票号",
                                   "借方金额", "贷方金额", "期间", "金额"], border, fill)
        for k, u in enumerate(self.result["unreached"], 3):
            r = u["record"]
            vals = [u["category"], r.source, r.date, r.settlement_method, r.ticket_no,
                    r.debit or None, r.credit or None,
                    "期初" if r.is_opening else "本期", r.amount()]
            for i, v in enumerate(vals, 1):
                c = ws.cell(row=k, column=i, value=v)
                c.border = border
                if i in (6, 7, 9) and isinstance(v, (int, float)):
                    c.number_format = "#,##0.00"
        for i, w in enumerate([20, 11, 12, 10, 12, 13, 13, 8, 13], 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # --------------------------------------------------------
    # 导出 CSV
    # --------------------------------------------------------
    def export_csv(self):
        if not self.result:
            messagebox.showinfo("提示", "请先点击【开始自动对账】生成结果后再导出。")
            return
        path = filedialog.asksaveasfilename(
            title="保存为 CSV", defaultextension=".csv",
            initialfile="银行对账结果_" + datetime.datetime.now().strftime("%Y%m%d_%H%M%S"),
            filetypes=[("CSV 文件", "*.csv")])
        if not path:
            return
        try:
            with open(path, "w", encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f)
                w.writerow(["银行对账系统 导出时间", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
                w.writerow([])
                w.writerow(["【一、全部流水清单】"])
                w.writerow(["期间", "来源", "日期", "结算方式", "票号", "借方金额", "贷方金额", "滚动余额", "对账标记"])
                for r in self.records:
                    w.writerow(["期初" if r.is_opening else "本期", r.source, r.date,
                                r.settlement_method, r.ticket_no,
                                r.debit or "", r.credit or "", r.balance or "", r.mark_label()])
                w.writerow([])
                w.writerow(["【二、已达账项】"])
                w.writerow(["企业日期", "企业票号", "企业金额", "银行日期", "银行票号", "银行金额", "匹配等级"])
                for er, br, level in self.result["pairs"]:
                    w.writerow([er.date, er.ticket_no, er.amount(),
                                br.date, br.ticket_no, br.amount(), level])
                w.writerow([])
                w.writerow(["【三、未达账项】"])
                w.writerow(["类别", "来源", "日期", "结算方式", "票号", "金额", "期间"])
                for u in self.result["unreached"]:
                    r = u["record"]
                    w.writerow([u["category"], r.source, r.date, r.settlement_method, r.ticket_no,
                                r.amount(), "期初" if r.is_opening else "本期"])
                w.writerow([])
                w.writerow(["【四、银行存款余额调节表】"])
                for ln in bc.statement_text(self.result["statement"]).splitlines():
                    w.writerow([ln])
        except Exception as e:
            messagebox.showerror("导出失败", str(e))
            return
        messagebox.showinfo("导出成功", "已保存至：\n{}".format(path))


# ============================================================
def main():
    root = tk.Tk()
    BankReconciliationApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
